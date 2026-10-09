"""SellerFlash Profit Report parsing — additive, fully separate from amazon_csv.

Supports .csv, .xlsx and legacy .xls files. Never mutates anything; returns
(parsed_rows, rejected_rows). Dates that cannot be parsed are kept as raw text
(reference only) instead of failing the row.
"""
import csv
import io
from datetime import date as calendar_date
from datetime import datetime
from decimal import Decimal, InvalidOperation

import xlrd
from openpyxl import load_workbook

from amazon_csv import normalize, parse_amount
from finance import money

SELLERFLASH_HEADERS = {
    "seller order id": "seller_order_id",
    "seller order status": "seller_order_status",
    "buyer order id": "buyer_order_id",
    "buyer order status": "buyer_order_status",
    "price": "price",
    "profit": "profit",
    "profit rate": "profit_rate",
    "order date": "order_date",
    "product cost": "product_cost",
    "initial shipment cost": "initial_shipment_cost",
    "seller refund": "seller_refund",
    "buyer refund": "buyer_refund",
    "marketplace": "marketplace",
    "delivery end date": "delivery_end_date",
    "notes": "notes",
}

NUMERIC_FIELDS = ("price", "profit", "profit_rate", "product_cost",
                  "initial_shipment_cost", "seller_refund", "buyer_refund")

MARKETPLACE_ALIASES = {
    "amazon.com": "US", "amazon.com.mx": "MX", "amazon.ca": "CA",
    "amazon.co.uk": "UK", "amazon.uk": "UK", "amazon.de": "DE",
    "amazon.fr": "FR", "amazon.it": "IT", "amazon.es": "ES",
    "amazon.nl": "NL", "amazon.se": "SE", "amazon.pl": "PL",
    "amazon.com.au": "AU", "amazon.au": "AU", "amazon.co.jp": "JP",
    "amazon.jp": "JP", "amazon.ae": "AE", "amazon.sa": "SA",
    "amazon.com.tr": "TR", "amazon.tr": "TR",
}

MAX_ROWS = 20000


def normalize_marketplace(value, valid_codes):
    """Map SellerFlash marketplace text (domain or code) to an internal code."""
    text = str(value or "").strip()
    if not text:
        return None
    upper = text.upper()
    if upper in valid_codes:
        return upper
    lowered = text.lower().replace("www.", "").strip().rstrip("/")
    lowered = lowered.split("/")[0].strip()
    return MARKETPLACE_ALIASES.get(lowered)


def _parse_number(value):
    """Tolerant numeric parse; returns float cents-rounded or None when empty."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(money(value))
    text = str(value).strip()
    if not text or text in ("-", "—"):
        return None
    cleaned = text.replace("%", "").strip()
    return float(parse_amount(cleaned))


def _parse_date(value):
    """Return (iso_date_or_None, raw_text). Dates are reference-only here."""
    if value is None:
        return None, ""
    if isinstance(value, datetime):
        return value.date().isoformat(), value.isoformat()
    if isinstance(value, calendar_date):
        return value.isoformat(), value.isoformat()
    text = str(value).strip()
    if not text:
        return None, ""
    head = text[:10]
    try:
        return calendar_date.fromisoformat(head).isoformat(), text
    except ValueError:
        pass
    for fmt in ("%b %d, %Y", "%d %b %Y", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat(), text
        except ValueError:
            continue
    return None, text


def _build_row(raw, line_number):
    """raw: SellerFlash alan adı → hücre değeri eşlemesi."""
    doc = {"line": line_number, **{field: raw.get(field) for field in SELLERFLASH_HEADERS.values()}}
    if not (str(doc.get("seller_order_id") or "").strip()):
        raise ValueError("Seller Order Id boş")
    doc["seller_order_id"] = str(doc["seller_order_id"]).strip()
    for field in NUMERIC_FIELDS:
        doc[field] = _parse_number(doc.get(field))
    doc["order_date"], doc["order_date_raw"] = _parse_date(doc.get("order_date"))
    doc["delivery_end_date"], doc["delivery_end_date_raw"] = _parse_date(doc.get("delivery_end_date"))
    for field in ("seller_order_status", "buyer_order_status", "buyer_order_id", "notes", "marketplace"):
        value = doc.get(field)
        doc[field] = "" if value is None else str(value).strip()
    return doc


def _rows_from_grid(grid):
    """grid: list of row-lists. Find the header row and yield (line, dict) pairs."""
    header_index = None
    mapping = {}
    for index, row in enumerate(grid[:15]):
        normalized = [normalize(cell) for cell in row]
        if "seller order id" in normalized:
            header_index = index
            mapping = {header: i for i, header in enumerate(normalized) if header}
            break
    if header_index is None:
        raise ValueError("SellerFlash başlıkları bulunamadı: 'Seller Order Id' sütunu gerekli")
    known = {field: mapping[header] for header, field in SELLERFLASH_HEADERS.items() if header in mapping}
    for index in range(header_index + 1, len(grid)):
        row = grid[index]
        if not any(str(cell or "").strip() for cell in row):
            continue
        yield index + 1, {field: (row[col] if col < len(row) else None) for field, col in known.items()}


def parse_sellerflash(content, filename):
    """Parse a SellerFlash report. Returns (accepted_rows, rejected_rows)."""
    name = (filename or "").lower()
    grid = []
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        sample = "\n".join(text.splitlines()[:5])
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        grid = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    elif name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif name.endswith(".xls"):
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        continue
                    except (ValueError, OverflowError):
                        pass
                row.append(cell.value)
            grid.append(row)
    else:
        raise ValueError("Desteklenmeyen dosya türü; .csv, .xlsx veya .xls yükleyin")

    accepted, rejected = [], []
    for line_number, raw in _rows_from_grid(grid):
        try:
            accepted.append(_build_row(raw, line_number))
        except (ValueError, InvalidOperation) as exc:
            rejected.append({"line": line_number, "reason": str(exc) or "Geçersiz değer"})
        if len(accepted) + len(rejected) > MAX_ROWS:
            raise ValueError(f"En fazla {MAX_ROWS} satırlık rapor yükleyin")
    return accepted, rejected
