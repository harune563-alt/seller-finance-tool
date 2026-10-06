import io
import os
import uuid
from datetime import date, timedelta

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient


# Company finance module regression: people/capital/debt/cash/closings/reporting/cron contracts
TEST_EMAIL = "company-test-76940be4f3@example.com"
TEST_PASSWORD = "T3st!Comp@ny2026#A"


def _read_frontend_base_url():
    env_path = "/app/frontend/.env"
    if not os.path.exists(env_path):
        return None
    with open(env_path, "r", encoding="utf-8") as fh:
        for line in fh:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    return None


def _read_backend_env_value(key: str):
    env_path = "/app/backend/.env"
    if not os.path.exists(env_path):
        return None
    with open(env_path, "r", encoding="utf-8") as fh:
        for line in fh:
            raw = line.strip()
            if not raw or raw.startswith("#") or "=" not in raw:
                continue
            k, v = raw.split("=", 1)
            if k == key:
                return v.strip().strip('"').strip("'")
    return None


def _prior_month_day(day: int = 10) -> str:
    first = date.today().replace(day=1)
    prev_last = first - timedelta(days=1)
    return prev_last.replace(day=min(day, 28)).isoformat()


@pytest.fixture(scope="session")
def base_url():
    value = os.environ.get("REACT_APP_BACKEND_URL") or _read_frontend_base_url()
    if not value:
        pytest.skip("REACT_APP_BACKEND_URL not configured")
    return value.rstrip("/")


@pytest.fixture(scope="session")
def api_session(base_url):
    s = requests.Session()
    login = s.post(
        f"{base_url}/api/auth/login",
        json={"email": TEST_EMAIL, "password": TEST_PASSWORD},
        timeout=30,
    )
    assert login.status_code == 200
    payload = login.json()
    assert isinstance(payload.get("token"), str) and payload["token"]
    s.headers.update({"Authorization": f"Bearer {payload['token']}"})
    return s


@pytest.fixture(scope="session")
def user_context(api_session, base_url):
    me = api_session.get(f"{base_url}/api/auth/me", timeout=30)
    assert me.status_code == 200
    data = me.json()
    assert data["email"] == TEST_EMAIL
    return data


@pytest.fixture(scope="session", autouse=True)
def cleanup_user_company_data(user_context):
    yield
    mongo_url = _read_backend_env_value("MONGO_URL")
    db_name = _read_backend_env_value("DB_NAME")
    if not mongo_url or not db_name:
        return
    uid = user_context["id"]
    client = MongoClient(mongo_url)
    db = client[db_name]
    db.company_jobs.delete_many({"user_id": uid})
    db.company_closings.delete_many({"user_id": uid})
    db.company_cash.delete_many({"user_id": uid})
    db.company_debts.delete_many({"user_id": uid})
    db.company_capital.delete_many({"user_id": uid})
    db.company_people.delete_many({"user_id": uid})
    db.transactions.delete_many({"user_id": uid})
    db.stores.delete_many({"user_id": uid})
    client.close()


def _create_store(api_session, base_url, name_suffix: str):
    body = {
        "name": f"TEST_COMPANY_{name_suffix}",
        "marketplaces": ["US"],
        "default_currency": "USD",
    }
    resp = api_session.post(f"{base_url}/api/stores", json=body, timeout=30)
    assert resp.status_code == 200
    return resp.json()


def _create_person(api_session, base_url, role: str, name: str):
    body = {"name": name, "role": role, "note": "test"}
    resp = api_session.post(f"{base_url}/api/company/people", json=body, timeout=30)
    assert resp.status_code == 200
    return resp.json()


def _capital(api_session, base_url, **kwargs):
    resp = api_session.post(f"{base_url}/api/company/capital", json=kwargs, timeout=30)
    return resp


def _tx_income(store_id: str, amount: float, dt: str, order_id: str, description: str = ""):
    return {
        "store_id": store_id,
        "marketplace": "US",
        "type": "income",
        "category": "Order payments",
        "amount": amount,
        "currency": "USD",
        "date": dt,
        "description": description,
        "order_id": order_id,
        "product_cost": 0,
        "shipping_cost": 0,
        "extra_cost": 0,
        "product_cost_recovery": 0,
        "shipping_cost_recovery": 0,
    }


def test_auth_cookie_and_me_for_test_user(base_url):
    login = requests.post(
        f"{base_url}/api/auth/login",
        json={"email": TEST_EMAIL, "password": TEST_PASSWORD},
        timeout=30,
    )
    assert login.status_code == 200
    cookie_header = login.headers.get("set-cookie", "")
    assert "access_token=" in cookie_header
    assert "httponly" in cookie_header.lower()
    token = login.json()["token"]
    me = requests.get(f"{base_url}/api/auth/me", headers={"Authorization": f"Bearer {token}"}, timeout=30)
    assert me.status_code == 200
    assert me.json()["email"] == TEST_EMAIL


def test_people_add_and_edit(api_session, base_url):
    partner = _create_person(api_session, base_url, "partner", f"TEST Partner {uuid.uuid4().hex[:6]}")
    assert partner["role"] == "partner"
    edit = api_session.patch(
        f"{base_url}/api/company/people/{partner['id']}",
        json={"name": partner["name"], "role": "contact", "note": "updated"},
        timeout=30,
    )
    assert edit.status_code == 200
    assert edit.json()["role"] == "contact"


def test_capital_ownership_math_and_withdrawal(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    person_a = _create_person(api_session, base_url, "partner", f"TEST A {uuid.uuid4().hex[:4]}")
    person_b = _create_person(api_session, base_url, "investor", f"TEST B {uuid.uuid4().hex[:4]}")
    dt = _prior_month_day(12)

    r1 = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person_a["id"],
        store_id=store["id"],
        currency="USD",
        direction="contribution",
        amount=3000,
        date=dt,
        note="A capital",
    )
    assert r1.status_code == 200
    r2 = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person_b["id"],
        store_id=store["id"],
        currency="USD",
        direction="contribution",
        amount=1000,
        date=dt,
        note="B capital",
    )
    assert r2.status_code == 200
    summary = r2.json()
    shares = {x["person_id"]: x["share_percent"] for x in summary["ownership"]}
    assert pytest.approx(shares[person_a["id"]], abs=0.01) == 75.0
    assert pytest.approx(shares[person_b["id"]], abs=0.01) == 25.0

    withdraw = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person_a["id"],
        store_id=store["id"],
        currency="USD",
        direction="withdrawal",
        amount=1000,
        date=dt,
        note="A withdraw",
    )
    assert withdraw.status_code == 200
    updated = withdraw.json()
    shares2 = {x["person_id"]: x["share_percent"] for x in updated["ownership"]}
    assert pytest.approx(shares2[person_a["id"]], abs=0.02) == 66.6667
    assert pytest.approx(shares2[person_b["id"]], abs=0.02) == 33.3333


def test_capital_idempotency_and_withdrawal_limits(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    person = _create_person(api_session, base_url, "partner", f"TEST C {uuid.uuid4().hex[:4]}")
    dt = _prior_month_day(9)
    req = str(uuid.uuid4())
    payload = {
        "request_id": req,
        "person_id": person["id"],
        "store_id": store["id"],
        "currency": "USD",
        "direction": "contribution",
        "amount": 100,
        "date": dt,
        "note": "same request",
    }
    first = _capital(api_session, base_url, **payload)
    assert first.status_code == 200
    second = _capital(api_session, base_url, **payload)
    assert second.status_code == 200
    assert len(second.json()["entries"]) == 1

    over = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person["id"],
        store_id=store["id"],
        currency="USD",
        direction="withdrawal",
        amount=200,
        date=dt,
        note="too much",
    )
    assert over.status_code == 422


def test_withdrawal_before_historical_funding_blocked(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    person = _create_person(api_session, base_url, "investor", f"TEST D {uuid.uuid4().hex[:4]}")
    contribution = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person["id"],
        store_id=store["id"],
        currency="USD",
        direction="contribution",
        amount=200,
        date=_prior_month_day(20),
        note="late funding",
    )
    assert contribution.status_code == 200
    bad = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person["id"],
        store_id=store["id"],
        currency="USD",
        direction="withdrawal",
        amount=10,
        date=_prior_month_day(1),
        note="before funding",
    )
    assert bad.status_code == 422


def test_debt_payment_flow_no_overpay_and_no_capital_effect(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    person = _create_person(api_session, base_url, "contact", f"TEST E {uuid.uuid4().hex[:4]}")
    dt = _prior_month_day(11)
    debt_req = str(uuid.uuid4())
    debt = api_session.post(
        f"{base_url}/api/company/debts",
        json={
            "request_id": debt_req,
            "person_id": person["id"],
            "store_id": store["id"],
            "direction": "payable",
            "currency": "USD",
            "amount": 500,
            "date": dt,
            "due_date": _prior_month_day(25),
            "cash_effect": True,
            "note": "debt create",
        },
        timeout=30,
    )
    assert debt.status_code == 200
    debt_id = debt.json()["id"]

    pay = api_session.post(
        f"{base_url}/api/company/debts/{debt_id}/payments",
        json={"request_id": str(uuid.uuid4()), "amount": 120, "date": dt, "note": "partial"},
        timeout=30,
    )
    assert pay.status_code == 200
    assert pay.json()["remaining"] == 380

    overpay = api_session.post(
        f"{base_url}/api/company/debts/{debt_id}/payments",
        json={"request_id": str(uuid.uuid4()), "amount": 1000, "date": dt, "note": "too much"},
        timeout=30,
    )
    assert overpay.status_code == 422

    cap = api_session.get(f"{base_url}/api/company/capital", params={"store_id": store["id"]}, timeout=30)
    assert cap.status_code == 200
    assert cap.json()["total_usd"] == 0


def test_overview_has_usd_try_and_cash_effects(api_session, base_url):
    person = _create_person(api_session, base_url, "contact", f"TEST F {uuid.uuid4().hex[:4]}")
    dt = _prior_month_day(10)

    cash_usd = api_session.post(
        f"{base_url}/api/company/cash",
        json={
            "request_id": str(uuid.uuid4()),
            "currency": "USD",
            "direction": "in",
            "amount": 75,
            "date": dt,
            "note": "open usd",
        },
        timeout=30,
    )
    assert cash_usd.status_code == 200

    cash_try = api_session.post(
        f"{base_url}/api/company/cash",
        json={
            "request_id": str(uuid.uuid4()),
            "currency": "TRY",
            "direction": "in",
            "amount": 50,
            "date": dt,
            "note": "open try",
        },
        timeout=30,
    )
    assert cash_try.status_code == 200

    debt_try = api_session.post(
        f"{base_url}/api/company/debts",
        json={
            "request_id": str(uuid.uuid4()),
            "person_id": person["id"],
            "store_id": None,
            "direction": "receivable",
            "currency": "TRY",
            "amount": 25,
            "date": dt,
            "cash_effect": False,
            "note": "try receivable",
        },
        timeout=30,
    )
    assert debt_try.status_code == 200

    overview = api_session.get(f"{base_url}/api/company/overview", timeout=30)
    assert overview.status_code == 200
    data = overview.json()
    by_currency = {x["currency"]: x for x in data["balances"]}
    assert "USD" in by_currency and "TRY" in by_currency
    assert by_currency["USD"]["cash_balance"] >= 75
    assert by_currency["TRY"]["cash_balance"] >= 50
    assert by_currency["TRY"]["receivables"] >= 25


def test_manual_closing_job_and_only_prior_month_included(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    previous_date = _prior_month_day(14)
    current_date = date.today().isoformat()
    sale_prev = api_session.post(
        f"{base_url}/api/transactions",
        json=_tx_income(store["id"], 100, previous_date, f"ORD-{uuid.uuid4().hex[:6]}"),
        timeout=30,
    )
    assert sale_prev.status_code == 200
    fee_prev = api_session.post(
        f"{base_url}/api/transactions",
        json={
            **_tx_income(store["id"], 30, previous_date, f"ORD-{uuid.uuid4().hex[:6]}"),
            "type": "expense",
            "category": "Service Fees",
        },
        timeout=30,
    )
    assert fee_prev.status_code == 200
    sale_current = api_session.post(
        f"{base_url}/api/transactions",
        json=_tx_income(store["id"], 60, current_date, f"ORD-{uuid.uuid4().hex[:6]}"),
        timeout=30,
    )
    assert sale_current.status_code == 200

    run = api_session.post(f"{base_url}/api/company/closings/run", timeout=30)
    assert run.status_code == 202
    assert run.json()["accepted"] is True

    completed = False
    for _ in range(20):
        jobs = api_session.get(f"{base_url}/api/company/jobs", timeout=30)
        assert jobs.status_code == 200
        items = jobs.json()
        if items and items[0]["status"] == "completed":
            completed = True
            break
    assert completed is True

    closings = api_session.get(f"{base_url}/api/company/closings", timeout=30)
    assert closings.status_code == 200
    rows = [c for c in closings.json() if c["store_id"] == store["id"]]
    assert len(rows) >= 1
    assert all(c["period"] == previous_date[:7] for c in rows)


def test_manual_closing_repeat_no_duplicates(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    previous_date = _prior_month_day(8)
    tx = api_session.post(
        f"{base_url}/api/transactions",
        json=_tx_income(store["id"], 110, previous_date, f"ORD-{uuid.uuid4().hex[:6]}"),
        timeout=30,
    )
    assert tx.status_code == 200

    first = api_session.post(f"{base_url}/api/company/closings/run", timeout=30)
    assert first.status_code == 202
    second = api_session.post(f"{base_url}/api/company/closings/run", timeout=30)
    assert second.status_code == 202

    closings = api_session.get(f"{base_url}/api/company/closings", timeout=30)
    assert closings.status_code == 200
    periods = [c for c in closings.json() if c["store_id"] == store["id"] and c["period"] == previous_date[:7]]
    assert len(periods) == 1


def test_store_delete_409_when_company_history_exists(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    person = _create_person(api_session, base_url, "partner", f"TEST G {uuid.uuid4().hex[:4]}")
    cap = _capital(
        api_session,
        base_url,
        request_id=str(uuid.uuid4()),
        person_id=person["id"],
        store_id=store["id"],
        currency="USD",
        direction="contribution",
        amount=50,
        date=_prior_month_day(9),
        note="history",
    )
    assert cap.status_code == 200
    delete_resp = api_session.delete(f"{base_url}/api/stores/{store['id']}", timeout=30)
    assert delete_resp.status_code == 409


def test_cron_endpoint_contracts(base_url):
    no_auth = requests.post(
        f"{base_url}/api/company/cron/monthly-close",
        json={"event": "schedule.triggered", "run_id": str(uuid.uuid4()), "data": None},
        timeout=30,
    )
    assert no_auth.status_code == 401

    bad_auth = requests.post(
        f"{base_url}/api/company/cron/monthly-close",
        headers={"Authorization": "Bearer wrong-secret"},
        json={"event": "schedule.triggered", "run_id": str(uuid.uuid4()), "data": None},
        timeout=30,
    )
    assert bad_auth.status_code == 401

    secret = _read_backend_env_value("WEBHOOK_CRON_SECRET")
    if not secret:
        pytest.skip("WEBHOOK_CRON_SECRET missing")
    bad_body = requests.post(
        f"{base_url}/api/company/cron/monthly-close",
        headers={"Authorization": f"Bearer {secret}"},
        json={"event": "invalid.event", "data": None},
        timeout=30,
    )
    assert bad_body.status_code == 400

    run_id = str(uuid.uuid4())
    ok = requests.post(
        f"{base_url}/api/company/cron/monthly-close",
        headers={"Authorization": f"Bearer {secret}", "X-Webhook-Id": run_id},
        json={"event": "schedule.triggered", "run_id": run_id, "data": None},
        timeout=30,
    )
    assert ok.status_code == 202
    assert ok.json()["accepted"] is True

    duplicate = requests.post(
        f"{base_url}/api/company/cron/monthly-close",
        headers={"Authorization": f"Bearer {secret}", "X-Webhook-Id": run_id},
        json={"event": "schedule.triggered", "run_id": run_id, "data": None},
        timeout=30,
    )
    assert duplicate.status_code == 202
    assert duplicate.json()["duplicate"] is True


def test_reports_orders_pagination_and_date_scope(api_session, base_url):
    store = _create_store(api_session, base_url, uuid.uuid4().hex[:6])
    in_range = _prior_month_day(15)
    out_range = (date.today().replace(day=1)).isoformat()

    for i in range(22):
        tx = api_session.post(
            f"{base_url}/api/transactions",
            json=_tx_income(store["id"], 10 + i, in_range, f"TEST-ORD-{i:03d}"),
            timeout=30,
        )
        assert tx.status_code == 200

    outside = api_session.post(
        f"{base_url}/api/transactions",
        json=_tx_income(store["id"], 99, out_range, "OUTSIDE-RANGE"),
        timeout=30,
    )
    assert outside.status_code == 200

    query = {
        "start_date": in_range,
        "end_date": in_range,
        "store_id": "ALL",
        "marketplace": "ALL",
        "currency": "ALL",
        "page": 1,
    }
    page1 = api_session.get(f"{base_url}/api/reports/orders", params=query, timeout=30)
    assert page1.status_code == 200
    p1 = page1.json()
    assert p1["total"] >= 22
    assert len(p1["items"]) == 20
    assert all(item["order_id"] != "OUTSIDE-RANGE" for item in p1["items"])

    page2 = api_session.get(f"{base_url}/api/reports/orders", params={**query, "page": 2}, timeout=30)
    assert page2.status_code == 200
    p2 = page2.json()
    assert p2["total_pages"] >= 2
    assert len(p2["items"]) >= 2


def test_excel_export_sheet_count_and_safety(api_session, base_url):
    store = _create_store(api_session, base_url, f"={uuid.uuid4().hex[:5]}")
    dt = _prior_month_day(16)
    order_id = "0000123456789012345"
    desc = "=INJECT"
    tx = api_session.post(
        f"{base_url}/api/transactions",
        json=_tx_income(store["id"], 123, dt, order_id, desc),
        timeout=30,
    )
    assert tx.status_code == 200

    excel = api_session.get(
        f"{base_url}/api/reports/excel",
        params={"start_date": dt, "end_date": dt, "store_id": "ALL", "marketplace": "ALL", "currency": "ALL"},
        timeout=60,
    )
    assert excel.status_code == 200
    assert excel.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    wb = load_workbook(io.BytesIO(excel.content), data_only=False)
    expected = ["Özet", "Pazar Yerleri", "Siparişler", "İşlem Detayları", "Amazon Ödemeleri", "Rapor Bilgisi"]
    assert wb.sheetnames == expected

    orders_sheet = wb["Siparişler"]
    values = [row for row in orders_sheet.iter_rows(min_row=2, values_only=True)]
    order_ids = [r[1] for r in values if r and r[1]]
    assert order_id in order_ids

    details_sheet = wb["İşlem Detayları"]
    detail_rows = [row for row in details_sheet.iter_rows(min_row=2, values_only=True)]
    descriptions = [r[17] for r in detail_rows if r and len(r) > 17]
    assert "'=INJECT" in descriptions
