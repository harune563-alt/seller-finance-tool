from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from finance import COST_FIELDS, RECOVERY_FIELDS, money

def safe(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def sheet(workbook, title, headers, rows):
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="047857")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in rows:
        ws.append([safe(value) for value in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(38, max(16, len(header) + 3))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00;[Red]-#,##0.00'
            cell.alignment = Alignment(vertical="top")
    return ws

def build_workbook(summary, rows, stores, orders):
    wb = Workbook(); wb.remove(wb.active)
    sheet(wb, "Özet", ["Mağaza", "Gelir USD", "Gider USD", "Net Kâr USD", "Marj %"],
          [[s["store_name"], s["revenue"], s["expenses"], s["net_profit"], s["margin"] if s["revenue"] > 0 else None] for s in summary["by_store"]] +
          [["TOPLAM", summary["revenue"], summary["expenses"], summary["net_profit"], summary["margin"] if summary["revenue"] > 0 else None]])
    sheet(wb, "Pazar Yerleri", ["Pazar", "Gelir USD", "Gider USD", "Net Kâr USD"], [[p["marketplace"], p["revenue"], p["expenses"], p["net"]] for p in summary["by_marketplace"]])
    sheet(wb, "Siparişler", ["Mağaza", "Order ID", "Pazar", "Orijinal Döviz", "İlk Tarih", "Son Tarih", "Hareket Sayısı", "Gelir USD", "Gider USD", "Net Kâr USD", "Marj %"],
          [[o["store_name"], o["order_id"] or "Siparişsiz işlem", o["marketplace"], o["currency"], o["first_date"], o["last_date"], o["records"], o["revenue"], o["expenses"], o["net_profit"], o["margin"] if o["revenue"] > 0 else None] for o in orders])
    details, payouts = [], []
    for tx in rows:
        store = stores.get(tx["store_id"], {}).get("name", "—")
        if tx["type"] == "payout":
            payouts.append([store, tx["date"], tx.get("payment_reference", ""), tx["marketplace"], tx["currency"], tx["amount"], tx["category"], tx.get("description", "")])
            continue
        income = tx["type"] == "income"
        costs = tx["usd_costs"]
        net = money(tx["amount_usd"]) * (1 if income else -1)
        if income: net -= sum(money(costs.get(key, 0)) for key in COST_FIELDS)
        if tx["category"] == "Refunds": net += sum(money(costs.get(key, 0)) for key in RECOVERY_FIELDS)
        details.append([store, tx["date"], tx.get("order_id", ""), tx["category"], tx["marketplace"], tx["currency"], tx["amount"] * (1 if income else -1), tx["amount_usd"] * (1 if income else -1),
                        *[costs.get(key, 0) for key in [*COST_FIELDS, *RECOVERY_FIELDS]], float(net), tx["fx"]["rate"], tx["fx"]["rate_date"], tx["fx"]["source"], tx.get("description", ""), tx["id"], tx.get("cost_currency", tx["currency"]),
                        *[tx.get(key, 0) for key in [*COST_FIELDS, *RECOVERY_FIELDS]]])
    sheet(wb, "İşlem Detayları", ["Mağaza", "Tarih", "Order ID", "Tür", "Pazar", "Orijinal Döviz", "Orijinal Tutar", "İşlem Tutarı USD", "Ürün Maliyeti USD", "Kargo Maliyeti USD", "Ekstra Maliyet USD", "Ürün Geri Kazanımı USD", "Kargo Geri Kazanımı USD", "Net Etki USD", "USD Kuru", "Kur Tarihi", "Kur Kaynağı", "Açıklama", "Kayıt ID", "Orijinal Maliyet Dövizi", "Orijinal Ürün Maliyeti", "Orijinal Kargo Maliyeti", "Orijinal Ekstra Maliyet", "Orijinal Ürün Geri Kazanımı", "Orijinal Kargo Geri Kazanımı"], details)
    sheet(wb, "Amazon Ödemeleri", ["Mağaza", "Tarih", "Ödeme Referansı", "Pazar", "Döviz", "Tutar", "Durum", "Not"], payouts)
    sheet(wb, "Rapor Bilgisi", ["Alan", "Değer"], [
        ["Başlangıç", summary["start_date"]], ["Bitiş", summary["end_date"]], ["Rapor dövizi", "USD"], ["Kaynak döviz filtresi", summary["source_currency"]],
        ["Kapsam", "Seçili tarih aralığındaki tüm hareketler; sipariş özetleri de bu dönem içindir."],
        ["Nakit ve kâr", "Amazon ödemeleri gelir değildir; kendi dövizinde ayrı sayfadadır. Sermaye/borç hareketleri mağaza P&L'sine dahil değildir."],
        ["Kur", "İşlem tarihine ait kayda sabitlenmiş Frankfurter referans kuru; gerçek banka kuru değildir."],
        ["Maliyetler", "Yeni maliyetler USD; eski maliyetlerin hem orijinali hem USD karşılığı İşlem Detayları sayfasındadır."],
    ])
    output = BytesIO(); wb.save(output); return output.getvalue()