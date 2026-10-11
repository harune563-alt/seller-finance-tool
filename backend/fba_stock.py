"""FBA lot CSV/XLSX import, per-SKU low-stock thresholds and stock status.
Lots only; never touches transactions or historical FIFO allocations."""
import csv
import io
import uuid
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from amazon_csv import normalize_header, parse_amount
from fba import OPEN, allocate_store, lot_costs, now_iso, q2
from sellerflash import _parse_date

HEADERS = {
    "sku": "sku", "asin": "asin", "productname": "product_name", "purchasedate": "purchase_date",
    "originalquantity": "original_quantity", "quantity": "original_quantity",
    "productunitcostusd": "product_unit_cost_usd", "productunitcost": "product_unit_cost_usd",
    "inboundshippingtotalusd": "inbound_shipping_total_usd", "inboundshippingtotal": "inbound_shipping_total_usd",
    "fbareceiveddate": "fba_received_date", "marketplace": "marketplace", "lotreference": "lot_reference",
}
REQUIRED = ("sku", "asin", "product_name", "purchase_date", "original_quantity",
            "product_unit_cost_usd", "inbound_shipping_total_usd")
COMPARE = ("asin", "product_name", "marketplace", "fba_received_date", "original_quantity",
           "product_unit_cost_usd", "inbound_shipping_total_usd")
LOCKED_WHEN_CONSUMED = ("original_quantity", "product_unit_cost_usd")
NEW, UNCHANGED, UPDATED, INVALID = "new_lot", "existing_unchanged", "updated", "invalid"
OUT, LOW, HEALTHY = "OUT_OF_STOCK", "LOW_STOCK", "HEALTHY"


class ThresholdIn(BaseModel):
    store_id: str
    sku: str = Field(min_length=1, max_length=200)
    threshold: Optional[int] = Field(default=None, ge=0, le=10_000_000)


def read_grid(content, filename):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig")
        try:
            delim = csv.Sniffer().sniff("\n".join(text.splitlines()[:5]), delimiters=",;\t").delimiter
        except csv.Error:
            delim = ","
        return list(csv.reader(io.StringIO(text), delimiter=delim))
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        grid = [list(r) for r in wb.active.iter_rows(values_only=True)]
        wb.close()
        return grid
    raise ValueError("Lot dosyası için .csv veya .xlsx seçin")


def parse_lot_rows(grid, valid_marketplaces):
    header_idx = next((i for i, r in enumerate(grid[:10])
                       if "sku" in [normalize_header(c) for c in r if c is not None]), None)
    if header_idx is None:
        raise ValueError("Başlık satırı bulunamadı (SKU sütunu gerekli)")
    cols = {HEADERS[h]: i for i, h in enumerate(normalize_header(c) if c is not None else "" for c in grid[header_idx])
            if h in HEADERS}
    missing = [f for f in REQUIRED if f not in cols]
    if missing:
        raise ValueError("Eksik sütunlar: " + ", ".join(missing))
    out = []
    for idx in range(header_idx + 1, len(grid)):
        raw = grid[idx]
        if not any(str(c or "").strip() for c in raw):
            continue
        cell = {f: (raw[i] if i < len(raw) else None) for f, i in cols.items()}
        text = {f: ("" if v is None else str(v).strip()) for f, v in cell.items()}
        row = {"line": idx + 1, "sku": text["sku"], "asin": text["asin"], "product_name": text["product_name"],
               "marketplace": text.get("marketplace", "").upper(), "lot_reference": text.get("lot_reference", ""),
               "errors": []}
        for f in ("purchase_date", "fba_received_date"):
            iso, _ = _parse_date(cell.get(f))
            row[f] = iso
            if f == "purchase_date" and not iso:
                row["errors"].append("Geçersiz alış tarihi")
            if f == "fba_received_date" and text.get(f) and not iso:
                row["errors"].append("Geçersiz FBA teslim tarihi")
        try:
            qty = parse_amount(text["original_quantity"])
            row["original_quantity"] = int(qty)
            if qty != int(qty) or qty <= 0:
                raise ValueError
        except Exception:
            row["original_quantity"] = None
            row["errors"].append("Adet pozitif tam sayı olmalı")
        for f in ("product_unit_cost_usd", "inbound_shipping_total_usd"):
            try:
                value = parse_amount(text[f])
                if value < 0:
                    raise ValueError
                row[f] = float(q2(value))
            except Exception:
                row[f] = None
                row["errors"].append(f"Geçersiz {f}")
        if not row["sku"]:
            row["errors"].append("SKU boş")
        if row["marketplace"] and row["marketplace"] not in valid_marketplaces:
            row["errors"].append("Pazar yeri mağazada aktif değil")
        out.append(row)
    return out


async def stock_status(db, user_id, store_id):
    lots = await db.fba_lots.find({"user_id": user_id, "store_id": store_id}, {"_id": 0}).to_list(None)
    thresholds = {t["sku"]: t.get("threshold") for t in
                  await db.fba_stock_thresholds.find({"user_id": user_id, "store_id": store_id}, {"_id": 0}).to_list(None)}
    by_sku = {}
    for lot in lots:
        b = by_sku.setdefault(lot["sku"], {"sku": lot["sku"], "asin": lot.get("asin") or "",
                                           "product_name": lot.get("product_name") or "", "remaining_units": 0, "open_lots": 0})
        if lot["lot_status"] == OPEN:
            b["remaining_units"] += lot["remaining_quantity"]
            b["open_lots"] += 1
    rows = []
    for sku, b in by_sku.items():
        t = thresholds.get(sku)
        status = OUT if b["remaining_units"] == 0 else LOW if t is not None and b["remaining_units"] <= t else HEALTHY
        rows.append({**b, "threshold": t, "stock_status": status})
    order = {OUT: 0, LOW: 1, HEALTHY: 2}
    return sorted(rows, key=lambda r: (order[r["stock_status"]], r["sku"]))


def fba_stock_router(db, get_current_user):
    router = APIRouter(prefix="/fba", tags=["fba"])

    async def own_store(store_id, user):
        store = await db.stores.find_one({"id": store_id, "user_id": user["id"]}, {"_id": 0})
        if not store:
            raise HTTPException(404, "Mağaza bulunamadı")
        return store

    @router.post("/lots/import")
    async def import_lots(store_id: str = Query(...), commit: bool = False, file: UploadFile = File(...),
                          user=Depends(get_current_user)):
        store = await own_store(store_id, user)
        content = await file.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(413, "Dosya en fazla 10 MB olabilir")
        try:
            rows = parse_lot_rows(read_grid(content, file.filename), set(store.get("marketplaces") or []))
        except (ValueError, UnicodeError) as exc:
            raise HTTPException(400, f"Lot dosyası okunamadı: {exc}")
        scope = {"user_id": user["id"], "store_id": store_id}
        seen, plan = set(), []
        for row in rows:
            key = (row["sku"], row["purchase_date"], row["lot_reference"])
            if key in seen:
                row["errors"].append("Dosyada aynı lot tekrar ediyor")
            seen.add(key)
            existing = None
            if not row["errors"]:
                ref_q = {"lot_reference": row["lot_reference"]} if row["lot_reference"] else \
                    {"lot_reference": {"$in": ["", None]}}
                # Stabil lot kimliği: Lot Reference varsa SKU + Lot Reference, yoksa SKU + Alış Tarihi
                q = {**scope, "sku": row["sku"], **ref_q}
                if not row["lot_reference"]:
                    q["purchase_date"] = row["purchase_date"]
                existing = await db.fba_lots.find_one(q, {"_id": 0}, sort=[("created_at", 1)])
            changes = {}
            if existing:
                changes = {f: [existing.get(f), row[f]] for f in COMPARE
                           if (row[f] or None) != (existing.get(f) or None) and not (f == "marketplace" and not row[f])}
                consumed = existing["remaining_quantity"] != existing["original_quantity"] or \
                    await db.fba_allocations.count_documents({"lot_id": existing["id"], "status": "active"})
                if consumed and any(f in changes for f in LOCKED_WHEN_CONSUMED):
                    row["errors"].append("Tüketilmiş lotta adet / ürün maliyeti değiştirilemez (FIFO geçmişi korunur)")
            status = INVALID if row["errors"] else UPDATED if changes else UNCHANGED if existing else NEW
            preview = {k: row.get(k) for k in ("line", "sku", "asin", "product_name", "purchase_date", "fba_received_date",
                                               "marketplace", "lot_reference", "original_quantity",
                                               "product_unit_cost_usd", "inbound_shipping_total_usd")}
            if status != INVALID:
                preview.update(lot_costs(row["original_quantity"], row["product_unit_cost_usd"], row["inbound_shipping_total_usd"]))
            plan.append({**preview, "status": status, "errors": row["errors"], "changes": changes,
                         "lot_code": (existing or {}).get("lot_code"), "_existing": existing})
        counts = {s: sum(1 for p in plan if p["status"] == s) for s in (NEW, UNCHANGED, UPDATED, INVALID)}
        if commit:
            seq = await db.fba_lots.count_documents(scope)
            ts = now_iso()
            for p in plan:
                ex = p["_existing"]
                fields = {f: p[f] for f in ("sku", "asin", "product_name", "purchase_date", "fba_received_date",
                                            "lot_reference", "original_quantity", "product_unit_cost_usd",
                                            "inbound_shipping_total_usd")}
                fields["marketplace"] = p["marketplace"] or (ex or {}).get("marketplace", "")
                if p["status"] == NEW:
                    seq += 1
                    lot = {"id": str(uuid.uuid4()), "lot_code": f"LOT-{seq:04d}", **scope, **fields,
                           **lot_costs(p["original_quantity"], p["product_unit_cost_usd"], p["inbound_shipping_total_usd"]),
                           "remaining_quantity": p["original_quantity"], "lot_status": OPEN, "source": "csv_import",
                           "created_at": ts, "updated_at": ts}
                    await db.fba_lots.insert_one(dict(lot))
                    await db.fba_lot_adjustments.insert_one({"id": str(uuid.uuid4()), **scope, "user_email": user.get("email", ""),
                        "lot_id": lot["id"], "lot_code": lot["lot_code"], "adjustment_type": "csv_import_create",
                        "previous_value": None, "new_value": {"original_quantity": lot["original_quantity"],
                        "landed_unit_cost_usd": lot["landed_unit_cost_usd"]}, "note": file.filename or "", "timestamp": ts})
                elif p["status"] == UPDATED:
                    sets = {**fields, **lot_costs(p["original_quantity"], p["product_unit_cost_usd"], p["inbound_shipping_total_usd"]),
                            "updated_at": ts}
                    if "original_quantity" in p["changes"]:
                        sets["remaining_quantity"] = p["original_quantity"]  # only reachable when not consumed
                    await db.fba_lots.update_one({"id": ex["id"]}, {"$set": sets})
                    await db.fba_lot_adjustments.insert_one({"id": str(uuid.uuid4()), **scope, "user_email": user.get("email", ""),
                        "lot_id": ex["id"], "lot_code": ex.get("lot_code", ""), "adjustment_type": "csv_import_update",
                        "previous_value": {f: c[0] for f, c in p["changes"].items()},
                        "new_value": {f: c[1] for f, c in p["changes"].items()}, "note": file.filename or "", "timestamp": ts})
            await allocate_store(db, user["id"], store_id)
        for p in plan:
            p.pop("_existing")
        return {"rows": plan, "counts": counts, "committed": commit}

    @router.get("/stock-status")
    async def get_stock_status(store_id: str, user=Depends(get_current_user)):
        await own_store(store_id, user)
        return await stock_status(db, user["id"], store_id)

    @router.put("/thresholds")
    async def set_threshold(data: ThresholdIn, user=Depends(get_current_user)):
        await own_store(data.store_id, user)
        await db.fba_stock_thresholds.update_one(
            {"user_id": user["id"], "store_id": data.store_id, "sku": data.sku.strip()},
            {"$set": {"threshold": data.threshold, "updated_at": now_iso()}}, upsert=True)
        return {"sku": data.sku.strip(), "threshold": data.threshold}

    return router
